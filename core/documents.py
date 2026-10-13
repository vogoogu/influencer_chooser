"""Read text from uploaded briefs (PDF, Word, plain text, markdown)."""
from __future__ import annotations

import io

MAX_CHARS = 20000


def extract_text(filename: str, data: bytes) -> str:
    name = filename.lower()
    if name.endswith((".txt", ".md", ".markdown")):
        text = data.decode("utf-8", errors="replace")
    elif name.endswith(".pdf"):
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        text = "\n".join((page.extract_text() or "") for page in reader.pages)
    elif name.endswith(".docx"):
        import docx
        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        text = "\n".join(parts)
    else:
        raise ValueError("Unsupported file type. Use .pdf, .docx, .txt or .md")
    text = text.strip()
    if not text:
        raise ValueError("No text found in the file (scanned PDF? paste the text instead).")
    return text[:MAX_CHARS]
